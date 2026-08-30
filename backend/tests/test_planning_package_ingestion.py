import io

from docx import Document
from openpyxl import Workbook

from app.services.planning_package_ingestion import detect_planning_resource_document


def workbook_bytes(headers,rows,title="Plan"):
    wb=Workbook();ws=wb.active;ws.title=title
    ws.append(headers)
    for row in rows:ws.append(row)
    stream=io.BytesIO();wb.save(stream)
    return stream.getvalue()


def test_staff_plan_is_detected_from_flexible_headers():
    content=workbook_bytes(
        ["Designation","Strength","Mobilization Date","Demobilization Date","Section"],
        [["Planning Manager",1,"2026-01-01","2027-12-31","Project"]],
        "Staff Deployment",
    )
    result=detect_planning_resource_document("Bidder Staff Plan.xlsx","xlsx",content)
    assert result["detected"] is True
    assert "Staff Plan" in result["plan_types"]


def test_equipment_plan_is_detected_without_fixed_template():
    content=workbook_bytes(
        ["Equipment","Qty","From","To","Activity"],
        [["Crane",2,"2026-01-01","2026-06-30","Mast Erection"]],
        "P&M",
    )
    result=detect_planning_resource_document("equipment_deployment.xlsx","xlsx",content)
    assert result["detected"] is True
    assert "Equipment Plan" in result["plan_types"]


def test_non_resource_spreadsheet_is_ignored():
    content=workbook_bytes(
        ["Item No","Description","Unit","Quantity"],
        [["1","OHE Mast","Nos",250]],
        "BOQ",
    )
    result=detect_planning_resource_document("BOQ.xlsx","xlsx",content)
    assert result["detected"] is False


def test_word_staff_plan_table_is_detected():
    document=Document()
    table=document.add_table(rows=1,cols=5)
    for i,value in enumerate(["Designation","Strength","Start Date","Finish Date","Section"]):
        table.rows[0].cells[i].text=value
    row=table.add_row().cells
    for i,value in enumerate(["Safety Manager","1","2026-01-01","2027-12-31","Project"]):
        row[i].text=value
    stream=io.BytesIO();document.save(stream)
    result=detect_planning_resource_document("Staff Deployment.docx","docx",stream.getvalue())
    assert result["detected"] is True
    assert "Staff Plan" in result["plan_types"]
